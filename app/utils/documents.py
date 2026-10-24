"""Documents: letters, invoices, free documents and CVs as PDF and / or Word file, written from what the user dictates.

The sums of an invoice are calculated here (never by the language model). The sender data comes from "My details" in the settings (routines.json, key 'profile').
Files go to  Documents\\Yume Dokumente.  PDF is made with Qt (no extra program), Word with python-docx.
"""
from __future__ import annotations

import html
import os
import re
from datetime import date, datetime
from pathlib import Path

DIR = Path(os.environ.get("USERPROFILE") or Path.home()) / "Documents" / "Yume Dokumente"


def profile() -> dict:
    from app.utils import routines
    return {"name": "", "street": "", "city": "", "email": "", "phone": "", "iban": "", "bic": "", "tax_id": "", "company": "", "kleinunternehmer": False,
            **routines.load().get("profile", {})}


def next_invoice_number() -> str:
    from app.utils import routines
    cfg = routines.load()
    p = cfg.setdefault("profile", {})
    n = int(p.get("invoice_counter", 0)) + 1
    p["invoice_counter"] = n
    routines.save(cfg)
    return f"R-{date.today().year}-{n:03d}"


def eur(x: float) -> str:
    return f"{x:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")


def _lines(value) -> list[str]:
    if isinstance(value, str):
        return [s.strip() for s in re.split(r"[\n;]+", value) if s.strip()]
    return [str(s).strip() for s in (value or []) if str(s).strip()]


def totals(items: list[dict], vat_percent: float) -> tuple[float, float, float]:
    net = sum(float(i.get("quantity", 1) or 0) * float(i.get("unit_price", 0) or 0) for i in items)
    vat = round(net * vat_percent / 100.0, 2)
    return round(net, 2), vat, round(net + vat, 2)


def _safe(name: str) -> str:
    return re.sub(r"[^\w\- äöüÄÖÜß]", "", name).strip().replace(" ", "_")[:60] or "Dokument"


def build(kind: str, f: dict) -> dict:
    """-> {'title', 'html', 'blocks'}: the document as HTML (for PDF) and as plain blocks (for Word)."""
    me = profile()
    today = f.get("date") or date.today().strftime("%d.%m.%Y")
    sender = [x for x in (me["company"] or me["name"], me["street"], me["city"]) if x]
    contact = [x for x in (("Tel. " + me["phone"]) if me["phone"] else "", me["email"]) if x]
    recipient = _lines(f.get("recipient"))
    blocks: list[tuple[str, object]] = []
    h = ["<html><body style='font-family:Arial; font-size:11pt; color:#111;'>"]

    def para(text, style=""):
        h.append(f"<p style='{style}'>{html.escape(text).replace(chr(10), '<br>')}</p>")
        blocks.append(("p", text))

    if kind == "brief":
        title = f.get("subject") or "Brief"
        h.append(f"<p style='font-size:9pt; color:#555;'>{html.escape(' · '.join(sender))}</p>")
        if recipient:
            h.append("<p>" + "<br>".join(html.escape(x) for x in recipient) + "</p>")
            blocks.append(("lines", recipient))
        h.append(f"<p style='text-align:right'>{html.escape((me['city'].split()[-1] + ', ') if me['city'] else '')}{html.escape(today)}</p>")
        blocks.append(("right", today))
        h.append(f"<p><b>{html.escape(title)}</b></p>")
        blocks.append(("bold", title))
        for p in _lines(f.get("body")) or [""]:
            para(p)
        para(f.get("closing") or "Mit freundlichen Grüßen")
        para(me["name"], "margin-top:30px;")
        return {"title": title, "html": "".join(h) + "</body></html>", "blocks": blocks}

    if kind == "rechnung":
        items = f.get("items") or []
        vat = 0.0 if me.get("kleinunternehmer") else float(f.get("vat_percent", 19) or 0)
        net, vat_amount, gross = totals(items, vat)
        number = f.get("invoice_number") or next_invoice_number()
        title = f"Rechnung {number}"
        h.append(f"<p style='font-size:9pt; color:#555;'>{html.escape(' · '.join(sender))}</p>")
        if recipient:
            h.append("<p>" + "<br>".join(html.escape(x) for x in recipient) + "</p>")
            blocks.append(("lines", recipient))
        h.append(f"<h2>{html.escape(title)}</h2><p>Datum: {html.escape(today)}" + (f"<br>Leistungszeitraum: {html.escape(str(f['period']))}" if f.get("period") else "") + "</p>")
        blocks.append(("bold", title)); blocks.append(("p", f"Datum: {today}"))
        h.append("<table width='100%' cellspacing='0' cellpadding='5' style='border-collapse:collapse;'><tr style='background:#eee;'><th align='left'>Beschreibung</th><th align='right'>Menge</th><th align='right'>Einzelpreis</th><th align='right'>Summe</th></tr>")
        rows = []
        for i in items:
            q, pr = float(i.get("quantity", 1) or 0), float(i.get("unit_price", 0) or 0)
            h.append(f"<tr><td style='border-bottom:1px solid #ddd'>{html.escape(str(i.get('description', '')))}</td><td align='right' style='border-bottom:1px solid #ddd'>{q:g}</td>"
                     f"<td align='right' style='border-bottom:1px solid #ddd'>{eur(pr)}</td><td align='right' style='border-bottom:1px solid #ddd'>{eur(q * pr)}</td></tr>")
            rows.append([str(i.get("description", "")), f"{q:g}", eur(pr), eur(q * pr)])
        h.append("</table>")
        blocks.append(("table", [["Beschreibung", "Menge", "Einzelpreis", "Summe"]] + rows))
        if me.get("kleinunternehmer"):
            h.append(f"<p align='right'><b>Gesamtbetrag: {eur(net)}</b></p><p style='font-size:9pt'>Kein Ausweis von Umsatzsteuer gemäß § 19 UStG (Kleinunternehmerregelung).</p>")
            blocks.append(("right", f"Gesamtbetrag: {eur(net)}")); blocks.append(("p", "Kein Ausweis von Umsatzsteuer gemäß § 19 UStG (Kleinunternehmerregelung)."))
        else:
            h.append(f"<p align='right'>Netto: {eur(net)}<br>Umsatzsteuer {vat:g} %: {eur(vat_amount)}<br><b>Gesamtbetrag: {eur(gross)}</b></p>")
            blocks.append(("right", f"Netto: {eur(net)}   USt {vat:g} %: {eur(vat_amount)}   Gesamtbetrag: {eur(gross)}"))
        pay = f.get("payment_terms") or "Bitte überweisen Sie den Betrag innerhalb von 14 Tagen ohne Abzug."
        para(pay)
        bank = " · ".join(x for x in (("IBAN " + me["iban"]) if me["iban"] else "", ("BIC " + me["bic"]) if me["bic"] else "", ("Steuernr. " + me["tax_id"]) if me["tax_id"] else "") if x)
        if bank:
            para(bank, "font-size:9pt; color:#555;")
        if f.get("notes"):
            para(str(f["notes"]))
        return {"title": title, "html": "".join(h) + "</body></html>", "blocks": blocks, "number": number, "gross": gross, "net": net}

    if kind == "lebenslauf":
        title = f"Lebenslauf {me['name']}".strip()
        h.append(f"<h1>{html.escape(me['name'] or 'Lebenslauf')}</h1><p>{html.escape(' · '.join(sender + contact))}</p>")
        blocks.append(("bold", me["name"] or "Lebenslauf")); blocks.append(("p", " · ".join(sender + contact)))
        for sec in f.get("sections") or []:
            h.append(f"<h3 style='border-bottom:1px solid #999'>{html.escape(str(sec.get('title', '')))}</h3>")
            blocks.append(("bold", str(sec.get("title", ""))))
            for it in _lines(sec.get("items")):
                h.append(f"<p style='margin:2px 0'>• {html.escape(it)}</p>")
                blocks.append(("p", "• " + it))
        return {"title": title, "html": "".join(h) + "</body></html>", "blocks": blocks}

    title = f.get("title") or f.get("subject") or "Dokument"
    h.append(f"<h1>{html.escape(title)}</h1>")
    blocks.append(("bold", title))
    for sec in f.get("sections") or []:
        if sec.get("title"):
            h.append(f"<h3>{html.escape(str(sec['title']))}</h3>")
            blocks.append(("bold", str(sec["title"])))
        for p in _lines(sec.get("text") or sec.get("items")):
            para(p)
    for p in _lines(f.get("body")):
        para(p)
    return {"title": title, "html": "".join(h) + "</body></html>", "blocks": blocks}


def write_pdf(doc: dict, path: Path):
    from PyQt6 import QtCore, QtGui
    from PyQt6.QtGui import QPageSize, QPdfWriter, QTextDocument
    writer = QPdfWriter(str(path))
    writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    writer.setPageMargins(QtCore.QMarginsF(22, 20, 20, 20), QtGui.QPageLayout.Unit.Millimeter)
    text = QTextDocument()
    text.setHtml(doc["html"])
    text.print(writer)


def write_docx(doc: dict, path: Path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    d = docx.Document()
    for kind, content in doc["blocks"]:
        if kind == "p":
            d.add_paragraph(str(content))
        elif kind == "bold":
            d.add_paragraph().add_run(str(content)).bold = True
        elif kind == "right":
            d.add_paragraph(str(content)).alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif kind == "lines":
            d.add_paragraph("\n".join(content))
        elif kind == "table":
            t = d.add_table(rows=0, cols=len(content[0]))
            t.style = "Table Grid"
            for row in content:
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
    d.save(str(path))


def create(kind: str, f: dict, fmt: str = "pdf") -> dict:
    kind = kind if kind in ("brief", "rechnung", "lebenslauf", "dokument") else "dokument"
    doc = build(kind, f)
    DIR.mkdir(parents=True, exist_ok=True)
    base = DIR / f"{datetime.now().strftime('%Y-%m-%d')}_{_safe(doc['title'])}"
    paths = []
    if fmt in ("pdf", "both"):
        write_pdf(doc, base.with_suffix(".pdf"))
        paths.append(base.with_suffix(".pdf"))
    if fmt in ("docx", "word", "both"):
        write_docx(doc, base.with_suffix(".docx"))
        paths.append(base.with_suffix(".docx"))
    if not paths:
        write_pdf(doc, base.with_suffix(".pdf"))
        paths.append(base.with_suffix(".pdf"))
    return {"title": doc["title"], "paths": [str(p) for p in paths], "number": doc.get("number"), "gross": doc.get("gross"), "net": doc.get("net")}


def recent(limit: int = 12) -> list[Path]:
    if not DIR.is_dir():
        return []
    return sorted([p for p in DIR.iterdir() if p.suffix.lower() in (".pdf", ".docx")], key=lambda p: p.stat().st_mtime, reverse=True)[:limit]
