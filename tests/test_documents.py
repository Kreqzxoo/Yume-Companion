"""Documents: letter, invoice (sums), CV, free document as PDF and Word; undo; tool plumbing. Temporary folder; Qt offscreen."""
import sys, os, tempfile, asyncio
os.environ["QT_QPA_PLATFORM"] = "offscreen"
sys.path.insert(0, '.')
from pathlib import Path
from PyQt6 import QtWidgets
app = QtWidgets.QApplication.instance() or QtWidgets.QApplication([])
from app.utils import documents as dc, routines, tool_log as tl
tmp = Path(tempfile.mkdtemp())
dc.DIR = tmp / "Yume Dokumente"; routines.FILE = tmp / "r.json"; tl.LOG_FILE = tmp / "tl.jsonl"; tl.enabled = lambda: True
cfg = routines.load(); cfg["profile"] = {"name": "Max Muster", "street": "Musterstr. 1", "city": "20095 Hamburg", "email": "max@example.com", "iban": "DE02120300000000202051", "kleinunternehmer": False}
routines.save(cfg)
ok = True
def check(name, cond):
    global ok
    ok &= bool(cond); print(("ok  " if cond else "BAD ") + name)
check("sums are calculated by the program", dc.totals([{"quantity": 2, "unit_price": 49.5}, {"quantity": 1, "unit_price": 100}], 19) == (199.0, 37.81, 236.81))
check("euro format", dc.eur(1234.5) == "1.234,50 €")
r = dc.create("brief", {"recipient": "Firma ABC; Hauptstr. 5; 10115 Berlin", "subject": "Kündigung meines Vertrags", "body": "hiermit kündige ich meinen Vertrag zum nächstmöglichen Zeitpunkt.\nBitte bestätigen Sie mir den Eingang.", "closing": "Freundliche Grüße"}, "both")
pdf = Path([p for p in r["paths"] if p.endswith(".pdf")][0]); docx_path = Path([p for p in r["paths"] if p.endswith(".docx")][0])
check("letter: PDF exists and is a PDF", pdf.is_file() and pdf.read_bytes()[:4] == b"%PDF" and pdf.stat().st_size > 1500)
import docx
text = "\n".join(p.text for p in docx.Document(str(docx_path)).paragraphs)
check("letter: Word file has subject, text and sender name", "Kündigung meines Vertrags" in text and "hiermit kündige" in text and "Max Muster" in text and "Firma ABC" in text)
inv = dc.create("rechnung", {"recipient": "Kunde GmbH; Weg 2; 50667 Köln", "items": [{"description": "Beratung", "quantity": 3, "unit_price": 80}, {"description": "Material", "quantity": 1, "unit_price": 40}], "period": "Oktober 2026"}, "both")
check("invoice: number and gross (280 + 19 % = 333,20)", inv["number"].startswith("R-") and abs(inv["gross"] - 333.20) < 0.005 and inv["number"].endswith("-001"))
d2 = docx.Document([p for p in inv["paths"] if p.endswith(".docx")][0])
check("invoice: Word table with the items", len(d2.tables) == 1 and d2.tables[0].rows[1].cells[0].text == "Beratung" and "240,00" in d2.tables[0].rows[1].cells[3].text)
inv2 = dc.create("rechnung", {"items": [{"description": "X", "quantity": 1, "unit_price": 10}]}, "pdf")
check("invoice numbers count up", inv2["number"].endswith("-002"))
cfg = routines.load(); cfg["profile"]["kleinunternehmer"] = True; routines.save(cfg)
k = dc.build("rechnung", {"items": [{"description": "X", "quantity": 1, "unit_price": 10}], "invoice_number": "R-1"})
check("small business: no VAT line, §19 note", "§ 19 UStG" in k["html"] and "Umsatzsteuer 19" not in k["html"])
cv = dc.create("lebenslauf", {"sections": [{"title": "Berufserfahrung", "items": ["2020-2024 Entwickler bei X", "2018-2020 Praktikant bei Y"]}, {"title": "Ausbildung", "items": "B.Sc. Informatik"}]}, "pdf")
check("CV pdf", Path(cv["paths"][0]).stat().st_size > 1500)
free = dc.create("dokument", {"title": "Packliste Rom", "sections": [{"title": "Kleidung", "items": ["Jacke", "Schirm"]}]}, "docx")
check("free document docx", "Packliste Rom" in "\n".join(p.text for p in docx.Document(free["paths"][0]).paragraphs))
from app.utils.ai_clients import tools as T
t = T.DocumentTool()
check("only 'open' needs approval", t.needs_approval({"action": "open"}) and not t.needs_approval({"action": "create"}))
res = asyncio.run(t.execute({"action": "create", "kind": "rechnung"}, {}))
check("invoice without items asks", not res["success"] and "Positionen" in res["result"])
res = asyncio.run(t.execute({"action": "create", "kind": "brief", "recipient": "A", "body": "Hallo"}, {}))
check("tool create letter + undo spec", res["success"] and res["undo"]["kind"] == "document_create")
tl.record("document", {"action": "create"}, "not_needed", True, "x", 0.1, undo=res["undo"])
check("undo deletes the file", tl.undo()[0] and not any(Path(p).exists() for p in res["undo"]["paths"]))
check("list", "pdf" in asyncio.run(t.execute({"action": "list"}, {}))["result"])
res = asyncio.run(t.execute({"action": "open", "path": "C:/Windows/notepad.exe"}, {}))
check("open refuses files outside the documents folder", not res["success"])
print("ALL OK" if ok else "FAILED")
